"""Read whatever a user uploads into clean all-text tables, or fail with a clear reason code.

Accepted:
* CSV / TSV / TXT: UTF-8, UTF-8 with BOM, UTF-16 (BOM) and Windows-1256 (common for Arabic
  exports); delimiter auto-detected among comma, semicolon, tab and pipe.
* Excel .xlsx / .xlsm: every non-empty sheet becomes a table.
* One file or several related files; relationships are found later by profiling.

Cleaning done here is structural only (never changes a value's meaning):
* header row detected; a file without one gets col_1..col_n;
* header names trimmed, invisible characters removed, blanks named, duplicates suffixed;
* fully empty rows and fully empty columns dropped (reported);
* every value kept as text (the engine infers types itself); Excel numbers that are whole
  become "42" not "42.0", dates become ISO.

Limits (documented in the UI and README): see LIMITS. Errors are IngestError(code, file) with
codes the web app maps to Arabic messages; messages never contain cell values.
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import PurePath

import pandas as pd

LIMITS = {
    "max_file_bytes": 15 * 1024 * 1024,
    "max_total_bytes": 30 * 1024 * 1024,
    "max_files": 10,
    "max_tables": 12,
    "max_rows_per_table": 200_000,
    "max_columns": 200,
}
CSV_EXT = {".csv", ".tsv", ".txt"}
EXCEL_EXT = {".xlsx", ".xlsm"}
_INVISIBLE = re.compile(r"[​-‏‪-‮⁦-⁩﻿ ]")
_NUMERIC = re.compile(r"^[+-]?[\d٠-٩۰-۹]+([.,٫][\d٠-٩۰-۹]+)?$")


class IngestError(ValueError):
    """code is stable (mapped to Arabic in the UI); file is the uploaded file name, never a value."""

    def __init__(self, code: str, file: str | None = None):
        super().__init__(code)
        self.code = code
        self.file = file


@dataclass
class TableNote:
    table: str
    source: str  # file name (and sheet)
    rows: int
    columns: int
    encoding: str | None = None
    delimiter: str | None = None
    header: str = "detected"  # detected | generated
    dropped_empty_rows: int = 0
    dropped_empty_columns: int = 0
    renamed_columns: int = 0
    skipped_title_rows: int = 0


@dataclass
class IngestResult:
    tables: dict[str, pd.DataFrame]
    notes: list[TableNote] = field(default_factory=list)

    def report(self) -> list[dict]:
        return [n.__dict__ for n in self.notes]


# ---------------------------------------------------------------- names

def clean_name(raw: object) -> str:
    text = "" if raw is None else str(raw)
    text = unicodedata.normalize("NFKC", _INVISIBLE.sub("", text))
    return re.sub(r"\s+", " ", text).strip()


def table_name(stem: str, taken: set[str]) -> str:
    base = re.sub(r"[^\w؀-ۿ]+", "_", clean_name(stem)).strip("_") or "table"
    base = base[:48]
    name, i = base, 2
    while name in taken:
        name, i = f"{base}_{i}", i + 1
    taken.add(name)
    return name


def _unique_headers(headers: list[str]) -> tuple[list[str], int]:
    out, seen, renamed = [], set(), 0
    for i, h in enumerate(headers, 1):
        name = clean_name(h) or f"col_{i}"
        if name != h:
            renamed += 1
        base, k = name, 2
        while name in seen:
            name, k = f"{base}_{k}", k + 1
            renamed += 1
        seen.add(name)
        out.append(name)
    return out, renamed


def _looks_like_header(first: list[str], rest: list[list[str]]) -> bool:
    """A header row is mostly non-empty, mostly non-numeric, distinct, and unlike the data below."""
    cells = [clean_name(c) for c in first]
    filled = [c for c in cells if c]
    if not filled or len(filled) < max(1, len(cells) // 2):
        return False
    numeric = sum(bool(_NUMERIC.match(c)) for c in filled)
    if numeric / len(filled) > 0.3:
        return False
    if not rest:
        return True
    # If the data below has numbers where the first row has words, the first row is a header.
    sample = rest[:20]
    for j, c in enumerate(cells):
        col = [r[j] for r in sample if j < len(r) and clean_name(r[j])]
        if col and sum(bool(_NUMERIC.match(clean_name(v))) for v in col) / len(col) > 0.6 and not _NUMERIC.match(c):
            return True
    # Long free-text cells in the first row are data, not names.
    if any(len(c) > 60 for c in filled):
        return False
    # Repeated first-row values in the data below mean the first row is data too.
    later = {clean_name(v) for r in sample for v in r}
    return not (set(filled) & later)


def _skip_title_rows(rows: list[list[str]]) -> tuple[list[list[str]], int]:
    """Drop report titles above the table: leading rows that fill far fewer cells than the rows below."""
    from collections import Counter

    filled = [sum(bool(clean_name(c)) for c in r) for r in rows[:60]]
    if len(filled) < 3:
        return rows, 0
    typical = Counter(filled[1:]).most_common(1)[0][0]
    if typical < 3:
        return rows, 0
    skip = 0
    while skip < min(5, len(rows) - 2) and filled[skip] <= max(1, typical // 3):
        skip += 1
    return rows[skip:], skip


def _frame(rows: list[list[str]], note: TableNote) -> pd.DataFrame:
    width = max((len(r) for r in rows), default=0)
    if width == 0:
        raise IngestError("empty_file", note.source)
    if width > LIMITS["max_columns"]:
        raise IngestError("too_many_columns", note.source)
    rows = [list(r) + [""] * (width - len(r)) for r in rows]
    rows, note.skipped_title_rows = _skip_title_rows(rows)
    if _looks_like_header(rows[0], rows[1:]):
        headers, data = rows[0], rows[1:]
    else:
        headers, data = [f"col_{i}" for i in range(1, width + 1)], rows
        note.header = "generated"
    headers, note.renamed_columns = _unique_headers(headers)
    df = pd.DataFrame(data, columns=headers, dtype=object)
    df = df.map(lambda v: None if v is None or (isinstance(v, str) and not clean_name(v)) else str(v).strip())
    before_rows, before_cols = len(df), df.shape[1]
    df = df.dropna(how="all")
    df = df.dropna(axis=1, how="all")
    note.dropped_empty_rows = before_rows - len(df)
    note.dropped_empty_columns = before_cols - df.shape[1]
    if df.empty or df.shape[1] == 0:
        raise IngestError("no_data", note.source)
    if len(df) > LIMITS["max_rows_per_table"]:
        raise IngestError("too_many_rows", note.source)
    df = df.reset_index(drop=True).astype(object).where(df.notna(), None)
    note.rows, note.columns = int(len(df)), int(df.shape[1])
    return df


# ---------------------------------------------------------------- CSV

def decode(data: bytes) -> tuple[str, str]:
    """(text, encoding). BOMs first, then strict UTF-8, then Windows-1256 (Arabic Windows)."""
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8", errors="replace"), "utf-8-sig"
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16"), "utf-16"
    try:
        return data.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        pass
    if b"\x00" in data[:4096]:
        raise IngestError("binary_file")
    return data.decode("cp1256", errors="replace"), "windows-1256"


def sniff_delimiter(text: str) -> str:
    """The delimiter that splits most lines into the same number of fields (a title line or two
    above the table does not change the answer). Ties prefer the larger field count."""
    from collections import Counter

    lines = [ln for ln in text.splitlines()[:200] if ln.strip()]
    if not lines:
        return ","
    best, best_score = ",", (-1.0, 0)
    for d in (",", ";", "	", "|"):
        counts = [len(next(csv.reader([ln], delimiter=d))) - 1 for ln in lines]
        mode, hits = Counter(c for c in counts if c > 0).most_common(1)[0] if any(counts) else (0, 0)
        if not mode:
            continue
        score = (hits / len(lines), mode)
        if score > best_score:
            best, best_score = d, score
    return best


def read_csv_bytes(data: bytes, source: str) -> tuple[pd.DataFrame, TableNote]:
    try:
        text, enc = decode(data)
    except IngestError as e:
        raise IngestError(e.code, source) from None
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    delim = sniff_delimiter(text)
    raw_rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    rows = [r for r in raw_rows if any(clean_name(c) for c in r)]
    note = TableNote(table="", source=source, rows=0, columns=0, encoding=enc,
                     delimiter={"\t": "tab"}.get(delim, delim))
    if not rows:
        raise IngestError("empty_file", source)
    df = _frame(rows, note)
    note.dropped_empty_rows += len(raw_rows) - len(rows)
    return df, note


# ---------------------------------------------------------------- Excel

def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == time(0) else v.isoformat(sep=" ", timespec="seconds")
    if isinstance(v, (date, time)):
        return v.isoformat()
    return str(v)


def read_excel_bytes(data: bytes, source: str) -> list[tuple[str, pd.DataFrame, TableNote]]:
    from openpyxl import load_workbook
    from zipfile import BadZipFile

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (BadZipFile, KeyError, OSError, ValueError):
        raise IngestError("bad_excel", source) from None
    out = []
    try:
        for ws in wb.worksheets:
            rows = []
            for r in ws.iter_rows(values_only=True):
                cells = [_cell(v) for v in r]
                if any(clean_name(c) for c in cells):
                    rows.append(cells)
                if len(rows) > LIMITS["max_rows_per_table"] + 1:
                    raise IngestError("too_many_rows", f"{source} / {ws.title}")
            if not rows:
                continue
            # trailing empty columns are common in Excel
            width = max(max((i + 1 for i, c in enumerate(r) if clean_name(c)), default=0) for r in rows)
            rows = [r[:width] for r in rows]
            note = TableNote(table="", source=f"{source} / {ws.title}", rows=0, columns=0)
            out.append((ws.title, _frame(rows, note), note))
    finally:
        wb.close()
    if not out:
        raise IngestError("no_data", source)
    return out


# ---------------------------------------------------------------- entry point

def read_files(files: list[tuple[str, bytes]]) -> IngestResult:
    """files: (original file name, content). Returns all-text tables and per-table notes."""
    if not files:
        raise IngestError("no_files")
    if len(files) > LIMITS["max_files"]:
        raise IngestError("too_many_files")
    if sum(len(b) for _, b in files) > LIMITS["max_total_bytes"]:
        raise IngestError("total_too_large")
    tables: dict[str, pd.DataFrame] = {}
    notes: list[TableNote] = []
    taken: set[str] = set()
    for name, data in files:
        stem, ext = PurePath(name).stem, PurePath(name).suffix.lower()
        if len(data) > LIMITS["max_file_bytes"]:
            raise IngestError("file_too_large", name)
        if not data.strip():
            raise IngestError("empty_file", name)
        if ext in EXCEL_EXT or data[:4] == b"PK\x03\x04":
            sheets = read_excel_bytes(data, name)
            for sheet, df, note in sheets:
                note.table = table_name(stem if len(sheets) == 1 else sheet, taken)
                tables[note.table] = df
                notes.append(note)
        elif ext == ".xls" or data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
            raise IngestError("old_excel", name)
        elif ext in CSV_EXT or ext == "":
            df, note = read_csv_bytes(data, name)
            note.table = table_name(stem, taken)
            tables[note.table] = df
            notes.append(note)
        else:
            raise IngestError("unsupported_type", name)
        if len(tables) > LIMITS["max_tables"]:
            raise IngestError("too_many_tables")
    return IngestResult(tables, notes)
