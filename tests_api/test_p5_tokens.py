"""P5 (verification token): recipients return ANY subset of rows; each row is verified by its
رمز_التحقق; admins re-link verified rows to the organization's own original data (no mapping table)."""
import io
import json
import logging
import random
import zipfile

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from nazeer import saudi_ids as s
from nazeer_api.models import AuditEvent, Organization, Return
from nazeer_api.security import encrypt_org_key, new_org_key
from nazeer_api.tokens import TOKEN_COLUMN
from tests_api.conftest import invite_and_join, signup
from tests_api.test_p2 import masked_twin, org_admin, ready_dataset, work  # noqa: F401

TOKEN_RE = r"^NZ-[A-Z2-7]+$"


# ---------------------------------------------------------------- data and helpers

def make_patients(n: int, seed: int = 1) -> pd.DataFrame:
    rng = random.Random(seed)
    return pd.DataFrame({
        "patient_id": [f"P-{i:06d}" for i in range(1, n + 1)],
        "full_name": [f"{s.gen_first_name(rng)} {s.gen_family_name(rng)}" for _ in range(n)],
        "national_id": [s.gen_saudi_id(rng, "1") for _ in range(n)],
        "mobile": [s.gen_mobile(rng) for _ in range(n)],
        "visit_date": [f"2026-{rng.randint(1, 12):02d}-{rng.randint(1, 28):02d}" for _ in range(n)],
        "amount": [str(rng.randint(100, 9000)) for _ in range(n)],
    })


def csv_bytes(df: pd.DataFrame) -> bytes:
    return df.to_csv(index=False).encode("utf-8")


def share_to_employee(app, admin, org, original: pd.DataFrame, email="employee@alwaha.example.com"):
    emp, _ = invite_and_join(app, admin, org, email)
    ds = ready_dataset(app, admin, org, [("patients.csv", csv_bytes(original))])
    twin = masked_twin(app, admin, org, ds["id"])
    assert twin["verdict"] == "PASS", twin["failed_checks"]
    member = next(m for m in admin.get(f"/api/orgs/{org}/members").json() if m["email"] == email)
    share = admin.post(f"/api/orgs/{org}/twins/{twin['id']}/shares", json={"member_ids": [member["id"]]}).json()
    return emp, share, twin, member


def download(client, share_id, fmt="xlsx") -> pd.DataFrame:
    url = client.post(f"/api/received/{share_id}/download", json={"format": fmt}).json()["url"]
    body = client.get(url).content
    if fmt == "xlsx":
        ws = load_workbook(io.BytesIO(body)).active
        rows = list(ws.iter_rows(values_only=True))
        return pd.DataFrame(rows[1:], columns=rows[0]).astype(object)
    with zipfile.ZipFile(io.BytesIO(body)) as z:
        name = z.namelist()[0]
        return pd.read_csv(z.open(name), dtype=str, keep_default_na=False, encoding="utf-8-sig")


def to_xlsx(df: pd.DataFrame) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def submit(client, share_id, data: bytes, name="results.xlsx"):
    return client.post(f"/api/received/{share_id}/returns", files=[("files", (name, data, "application/octet-stream"))])


def relink(app, admin, org, return_id, original: pd.DataFrame, confirm=False, name="patients.csv"):
    r = admin.post(f"/api/orgs/{org}/returns/{return_id}/relink", data={"confirm_partial": "true" if confirm else "false"},
                   files=[("files", (name, csv_bytes(original), "text/csv"))])
    if r.status_code != 202:
        return r
    work(app)
    return admin.get(f"/api/orgs/{org}/returns/{return_id}")


def relinked(admin, org, return_id) -> pd.DataFrame:
    r = admin.get(f"/api/orgs/{org}/returns/{return_id}/relinked.csv")
    assert r.status_code == 200, r.text
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        return pd.read_csv(z.open(z.namelist()[0]), dtype=str, keep_default_na=False, encoding="utf-8-sig")


def relinked_after(app, admin, org, return_id, original):
    d = relink(app, admin, org, return_id, original).json()
    assert d["relink"]["status"] == "ready", d["relink"]
    return relinked(admin, org, return_id)


@pytest.fixture(scope="module")
def patients():
    return make_patients(400)


# ---------------------------------------------------------------- the main story: a 1% subset

def test_return_100_of_10000_rows_verifies_and_relinks(app, org_admin):
    admin, org = org_admin
    original = make_patients(10_000, seed=7)
    emp, share, twin, _ = share_to_employee(app, admin, org, original)
    shared = download(emp, share["id"])
    assert len(shared) == 10_000 and shared[TOKEN_COLUMN].str.match(TOKEN_RE).all()
    assert list(shared.columns)[0] == TOKEN_COLUMN
    subset = shared.sample(100, random_state=3).copy()
    subset["risk_score"] = [f"{i / 100:.2f}" for i in range(100)]
    r = submit(emp, share["id"], to_xlsx(subset))
    assert r.status_code == 201, r.text
    rep = r.json()["report"]
    assert rep["counts"]["verified"] == 100 and rep["integrity"] == 1.0 and rep["coverage"] == 0.01
    assert rep["added_columns"] == {"patients": ["risk_score"]}

    out = relinked_after(app, admin, org, r.json()["id"], original)
    assert set(out["حالة_الربط"]) == {"مرتبط"} and len(out) == 100
    # the twin's patient_id was pseudonymized; the output carries the organization's real rows
    assert set(out["patient_id"]) <= set(original["patient_id"])
    by_id = original.set_index("patient_id")
    assert (out["national_id"] == by_id.loc[out["patient_id"], "national_id"].values).all()


# ---------------------------------------------------------------- the recipient's copy is never trusted

def test_edits_in_original_columns_are_ignored_and_counted(app, org_admin, patients):
    admin, org = org_admin
    emp, share, _, _ = share_to_employee(app, admin, org, patients)
    shared = download(emp, share["id"])
    sub = shared.head(50).copy()
    sub["score"] = [str(i) for i in range(50)]
    sub.loc[0, "national_id"], sub.loc[1, "national_id"] = sub.loc[1, "national_id"], sub.loc[0, "national_id"]  # swap IDs
    sub.loc[2, "full_name"] = "اسم معدّل"
    sub.loc[3, "mobile"] = "0500000000"
    sub.loc[4, "visit_date"] = "1999-01-01"
    r = submit(emp, share["id"], to_xlsx(sub)).json()
    assert r["report"]["counts"]["verified"] == 50
    assert r["report"]["rows_changed_in_twin_columns"] == 5          # informational only

    out = relinked_after(app, admin, org, r["id"], patients)
    real = patients.set_index("patient_id")
    for _, row in out.iterrows():                                      # true originals, whatever was edited
        assert row["national_id"] == real.loc[row["patient_id"], "national_id"]
        assert row["full_name"] == real.loc[row["patient_id"], "full_name"]
        assert row["mobile"] == real.loc[row["patient_id"], "mobile"]
    assert out["score"].tolist() == [str(i) for i in range(50)]       # recipient's added column, in order


# ---------------------------------------------------------------- Excel round trips

def test_excel_round_trips_keep_every_token(app, org_admin, patients):
    admin, org = org_admin
    emp, share, _, _ = share_to_employee(app, admin, org, patients)
    # 1) openpyxl: open the downloaded workbook, save it, return it
    url = emp.post(f"/api/received/{share['id']}/download", json={"format": "xlsx"}).json()["url"]
    wb = load_workbook(io.BytesIO(emp.get(url).content))
    cell = wb.active.cell(row=2, column=1)
    assert cell.data_type == "s" and str(cell.value).startswith("NZ-")      # written as text
    buf = io.BytesIO()
    wb.save(buf)
    r = submit(emp, share["id"], buf.getvalue()).json()
    assert r["report"]["integrity"] == 1.0 and r["report"]["counts"]["verified"] == len(patients)

    # 2) CSV opened and saved by Excel: number-like text becomes numbers; tokens stay text
    shared = download(emp, share["id"], "csv")
    excelled = shared.map(lambda v: str(int(v)) if isinstance(v, str) and v.isdigit() else v)
    r = submit(emp, share["id"], excelled.to_csv(index=False).encode("utf-8"), name="r.csv").json()
    assert r["report"]["integrity"] == 1.0


# ---------------------------------------------------------------- rejections and per-row statuses

def test_missing_token_column_cannot_relink(app, org_admin, patients):
    admin, org = org_admin
    emp, share, _, _ = share_to_employee(app, admin, org, patients)
    shared = download(emp, share["id"]).drop(columns=[TOKEN_COLUMN])
    r = submit(emp, share["id"], to_xlsx(shared))
    assert r.status_code == 422 and r.json()["code"] == "token_column_missing"


def test_duplicate_altered_missing_and_foreign_tokens(app, org_admin, patients):
    admin, org = org_admin
    emp, share, twin, member = share_to_employee(app, admin, org, patients)
    other = admin.post(f"/api/orgs/{org}/twins/{twin['id']}/shares", json={"member_ids": [member["id"]]}).json()
    mine, theirs = download(emp, share["id"]).head(20).copy(), download(emp, other["id"]).iloc[30:32].copy()
    t = mine.loc[5, TOKEN_COLUMN]
    mine.loc[5, TOKEN_COLUMN] = t[:-1] + ("A" if t[-1] != "A" else "B")       # one character changed
    mine.loc[6, TOKEN_COLUMN] = None                                            # token removed
    rows = pd.concat([mine, mine.iloc[[0]], theirs], ignore_index=True)         # a duplicate + 2 foreign
    r = submit(emp, share["id"], to_xlsx(rows)).json()["report"]
    c = r["counts"]
    assert c == {"verified": 18, "invalid": 1, "missing": 1, "foreign": 2, "old_key": 0, "duplicate": 1}
    assert r["rows_by_status"]["invalid"] == [6] and r["rows_by_status"]["missing"] == [7]
    assert r["rows_by_status"]["duplicate"] == [21]
    assert r["integrity"] == round(18 / 23, 4)

    # below 100%: the admin must confirm, and the confirmation is audited
    rid = admin.get(f"/api/orgs/{org}/returns").json()[0]["id"]
    resp = relink(app, admin, org, rid, patients)
    assert resp.status_code == 409 and resp.json()["code"] == "relink_confirm_partial"
    d = relink(app, admin, org, rid, patients, confirm=True).json()
    assert d["relink"]["matched"] == 18
    with app.state.sessionmaker() as db:
        assert db.execute(select(AuditEvent).where(AuditEvent.action == "relink.partial_confirmed")).first()


def test_token_from_another_organization_is_foreign(app, org_admin, patients):
    admin, org = org_admin
    emp, share, _, _ = share_to_employee(app, admin, org, patients)
    other_admin, me = signup(app, "boss@other.example.com", org="منشأة أخرى")
    other_org = me["memberships"][0]["org_id"]
    emp2, share2, _, _ = share_to_employee(app, other_admin, other_org, patients, email="x@other.example.com")
    foreign = download(emp2, share2["id"]).head(10)
    mine = download(emp, share["id"]).head(10)
    r = submit(emp, share["id"], to_xlsx(pd.concat([mine, foreign]))).json()["report"]
    assert r["counts"]["verified"] == 10 and r["counts"]["foreign"] == 10


def test_superset_or_subset_original_and_added_columns(app, org_admin, patients):
    admin, org = org_admin
    emp, share, _, _ = share_to_employee(app, admin, org, patients)
    sub = download(emp, share["id"]).iloc[100:130].copy()
    sub["segment"], sub["note"] = "A", "تمت المراجعة"
    rid = submit(emp, share["id"], to_xlsx(sub)).json()["id"]
    extra = make_patients(50, seed=99)
    extra["patient_id"] = "Q" + extra["patient_id"]
    bigger = pd.concat([patients, extra]).sample(frac=1, random_state=1)                 # superset, shuffled
    out = relinked_after(app, admin, org, rid, bigger)
    assert len(out) == 30 and set(out["حالة_الربط"]) == {"مرتبط"}
    assert {"segment", "note"} <= set(out.columns)
    assert list(out.columns)[:6] == list(patients.columns)                                # original rows as uploaded
    # a subset of the original lacking some returned rows: those link, the others are marked
    out2 = relinked_after(app, admin, org, rid, patients.iloc[:115])
    assert (out2["حالة_الربط"] == "مرتبط").sum() == 15


def test_key_rotation_reports_old_tokens(app, org_admin, patients):
    admin, org = org_admin
    emp, share, _, _ = share_to_employee(app, admin, org, patients)
    rows = download(emp, share["id"]).head(10)
    with app.state.sessionmaker() as db:
        o = db.get(Organization, org)
        o.key_version += 1
        o.enc_key = encrypt_org_key(app.state.settings.master_key, o.id, o.key_version, new_org_key())
        db.commit()
    r = submit(emp, share["id"], to_xlsx(rows)).json()["report"]
    assert r["counts"]["old_key"] == 10 and r["counts"]["verified"] == 0


def test_no_values_in_report_audit_or_logs(app, org_admin, patients, caplog):
    admin, org = org_admin
    caplog.set_level(logging.DEBUG)
    emp, share, _, _ = share_to_employee(app, admin, org, patients)
    sub = download(emp, share["id"]).head(40).copy()
    sub["score"] = "0.731"
    rid = submit(emp, share["id"], to_xlsx(sub)).json()["id"]
    relinked_after(app, admin, org, rid, patients)
    with app.state.sessionmaker() as db:
        metas = [e.meta for e in db.execute(select(AuditEvent)).scalars()]
        reports = [r.report for r in db.execute(select(Return)).scalars()]
    blob = json.dumps([metas, reports], ensure_ascii=False) + caplog.text
    secrets = set(patients["national_id"].head(40)) | set(patients["full_name"].head(40)) | set(patients["mobile"].head(40))
    secrets |= set(sub[TOKEN_COLUMN]) | set(patients["patient_id"].head(40)) | {"0.731"}
    assert not any(v in blob for v in secrets)


def test_roles_and_tenancy(app, org_admin, patients):
    admin, org = org_admin
    emp, share, _, _ = share_to_employee(app, admin, org, patients)
    rid = submit(emp, share["id"], to_xlsx(download(emp, share["id"]).head(5))).json()["id"]
    dm, _ = invite_and_join(app, admin, org, "dm@alwaha.example.com", data_manager=True)
    assert dm.get(f"/api/orgs/{org}/returns").status_code == 200
    r = dm.post(f"/api/orgs/{org}/returns/{rid}/relink", files=[("files", ("p.csv", b"a\n1", "text/csv"))])
    assert r.status_code == 403
    second, _ = invite_and_join(app, admin, org, "admin2@alwaha.example.com", role="admin")
    relinked_after(app, admin, org, rid, patients)
    assert second.get(f"/api/orgs/{org}/returns/{rid}/relinked.csv").json()["code"] == "relink_not_yours"
    other, _ = signup(app, "z@other.example.com", org="أخرى")
    assert other.get(f"/api/orgs/{org}/returns/{rid}").status_code == 404
    assert submit(other, share["id"], b"x").status_code == 404
