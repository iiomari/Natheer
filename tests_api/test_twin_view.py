"""Twin viewer (organization and recipients) and the one-page PDF report."""
import io

from openpyxl import load_workbook

from tests_api.test_p2 import demo_files, masked_twin, org_admin, ready_dataset  # noqa: F401
from tests_api.test_p5_tokens import make_patients, share_to_employee
from nazeer_api.tokens import TOKEN_COLUMN


def test_twin_view_pages_search_sort_and_marks(app, org_admin, demo_files):
    admin, org = org_admin
    ds = ready_dataset(app, admin, org, demo_files)
    twin = masked_twin(app, admin, org, ds["id"])
    base = f"/api/orgs/{org}/twins/{twin['id']}/rows"
    r = admin.get(base, params={"table": "customers", "size": 25}).json()
    assert r["table"] == "customers" and r["total"] == 150 and r["pages"] == 6 and len(r["rows"]) == 25
    assert r["columns"][0] == TOKEN_COLUMN and r["rows"][0]["cells"][TOKEN_COLUMN][0].startswith("NZ-")
    assert "national_id" in r["replaced_columns"]
    assert r["rows"][0]["cells"]["national_id"][1] == "replaced"
    assert r["rows"][0]["cells"]["city"][1] is None
    last = admin.get(base, params={"table": "customers", "size": 25, "page": 99}).json()
    assert last["page"] == 6 and len(last["rows"]) == 0 or len(last["rows"]) <= 25
    # free-text cells: replaced ones are marked
    claims = admin.get(base, params={"table": "claims", "size": 100}).json()
    assert any(row["cells"]["notes"][1] == "replaced" for row in claims["rows"])
    # search and numeric sort
    city = r["rows"][0]["cells"]["city"][0]
    found = admin.get(base, params={"table": "customers", "q": city, "size": 100}).json()
    assert found["total"] >= 1 and all(city in row["cells"]["city"][0] for row in found["rows"])
    ages = admin.get(base, params={"table": "customers", "sort": "age", "desc": "true", "size": 100}).json()
    vals = [row["cells"]["age"][0] for row in ages["rows"] if row["cells"]["age"][0]]
    assert vals == sorted(vals, key=lambda v: float(v.replace(",", "")), reverse=True)
    # downloads for the organization: Excel with the token column as text
    wb = load_workbook(io.BytesIO(admin.get(f"/api/orgs/{org}/twins/{twin['id']}/download.xlsx").content))
    assert wb["customers"].cell(row=1, column=1).value == TOKEN_COLUMN


def test_recipient_view_and_review_marks(app, org_admin):
    admin, org = org_admin
    original = make_patients(120)
    phrases = ["تمت المتابعة مع المريض في العيادة", "يحتاج المريض إلى تحاليل إضافية قبل الموعد القادم",
               "تم صرف العلاج حسب وصفة الطبيب المعالج", "المريض راجع قسم الطوارئ مساء أمس", "لا توجد ملاحظات إضافية"]
    original["notes"] = ["رقم الطلب 1000000008 لدى المختبر" if i == 3 else f"{phrases[i % 5]} رقم {i}" for i in range(120)]
    emp, share, twin, _ = share_to_employee(app, admin, org, original)
    r = emp.get(f"/api/received/{share['id']}/rows", params={"size": 10}).json()
    assert r["total"] == 120 and r["columns"][0] == TOKEN_COLUMN
    org_tok = admin.get(f"/api/orgs/{org}/twins/{twin['id']}/rows", params={"size": 10}).json()
    assert r["rows"][0]["cells"][TOKEN_COLUMN][0] != org_tok["rows"][0]["cells"][TOKEN_COLUMN][0]  # per-share tokens
    row4 = emp.get(f"/api/received/{share['id']}/rows", params={"size": 10, "q": "رقم الطلب"}).json()["rows"][0]
    assert row4["n"] == 4 and row4["cells"]["notes"][1] == "review"
    other, _ = __import__("tests_api.conftest", fromlist=["signup"]).signup(app, "o@x.example.com", org="أخرى")
    assert other.get(f"/api/received/{share['id']}/rows").status_code == 404


def test_pdf_report_is_one_page_arabic_and_value_free(app, org_admin, demo_files):
    from pypdf import PdfReader

    admin, org = org_admin
    ds = ready_dataset(app, admin, org, demo_files)
    twin = masked_twin(app, admin, org, ds["id"])
    r = admin.get(f"/api/orgs/{org}/twins/{twin['id']}/report.pdf")
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf" and r.content[:4] == b"%PDF"
    doc = PdfReader(io.BytesIO(r.content))
    assert len(doc.pages) == 1
    text = doc.pages[0].extract_text()
    # shaped Arabic comes out as presentation forms in visual order: check it is there, not substrings
    assert sum(0xFE70 <= ord(ch) <= 0xFEFF or 0x0600 <= ord(ch) <= 0x06FF for ch in text) > 200
    assert "demo" in text
    fonts = str(doc.pages[0]["/Resources"]["/Font"].get_object())
    assert "Plex" in str([f.get_object()["/BaseFont"] for f in doc.pages[0]["/Resources"]["/Font"].get_object().values()]) or "Plex" in fonts
    customers = dict(demo_files)["customers.csv"].decode("utf-8").splitlines()[1:20]
    ids = [line.split(",")[2] for line in customers]
    assert not any(v in text for v in ids)
