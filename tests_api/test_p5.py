"""P5: recipients return results; admins re-link them to the real records (in memory, no mapping table)."""
import io
import json
import logging
import re
import zipfile
from datetime import timedelta

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select, update

from nazeer.transform import Pseudonymizer
from nazeer_api.db import utcnow
from nazeer_api.models import AuditEvent, Blob, Organization, Return
from nazeer_api.processing import sweep
from nazeer_api.returns import REF_COLUMN, ReturnError, check_return, make_ref
from nazeer_api.storage import export_csv_zip, export_xlsx
from tests_api.conftest import invite_and_join
from tests_api.test_p2 import demo_files, masked_twin, org_admin, ready_dataset, work  # noqa: F401

REF_RE = re.compile(r"^NZ-[A-Z2-7]{12}$")


# ---------------------------------------------------------------- helpers

def _shared_with_employee(app, admin, org, files, email="employee@alwaha.example.com"):
    emp, _ = invite_and_join(app, admin, org, email)
    ds = ready_dataset(app, admin, org, files)
    twin = masked_twin(app, admin, org, ds["id"], apply_fix="auto")
    assert twin["verdict"] == "PASS", twin["verdict"]
    member = next(m for m in admin.get(f"/api/orgs/{org}/members").json() if m["email"] == email)
    share = admin.post(f"/api/orgs/{org}/twins/{twin['id']}/shares", json={"member_ids": [member["id"]]}).json()
    return emp, share, ds, twin


def _download_csv_tables(client, share_id) -> dict[str, pd.DataFrame]:
    url = client.post(f"/api/received/{share_id}/download", json={"format": "csv"}).json()["url"]
    out = {}
    with zipfile.ZipFile(io.BytesIO(client.get(url).content)) as z:
        for name in z.namelist():
            out[name[:-4]] = pd.read_csv(z.open(name), dtype=str, keep_default_na=False, encoding="utf-8-sig")
    return out


def _results_csv(df: pd.DataFrame) -> bytes:
    return df.to_csv(index=False).encode("utf-8")


def _submit(client, share_id, data: bytes, name="results.csv"):
    return client.post(f"/api/received/{share_id}/returns", files=[("files", (name, data, "text/csv"))])


def _relink(app, admin, org, return_id, files):
    r = admin.post(f"/api/orgs/{org}/returns/{return_id}/relink",
                   files=[("files", (n, d, "text/csv")) for n, d in files])
    assert r.status_code == 202, r.text
    work(app)
    return admin.get(f"/api/orgs/{org}/returns/{return_id}").json()


def _relinked_csv(admin, org, return_id) -> pd.DataFrame:
    r = admin.get(f"/api/orgs/{org}/returns/{return_id}/relinked.csv")
    assert r.status_code == 200, r.text
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        return pd.read_csv(z.open("results.csv"), dtype=str, keep_default_na=False, encoding="utf-8-sig")


def _scored(customers: pd.DataFrame) -> pd.DataFrame:
    """What a recipient sends back: the key, its reference, and a score (here: the row number)."""
    return pd.DataFrame({"customer_id": customers["customer_id"], REF_COLUMN: customers[REF_COLUMN],
                         "score": [str(i) for i in range(len(customers))]})


# ---------------------------------------------------------------- the full story

def test_return_and_relink_full_story(app, org_admin, demo_files, caplog):
    admin, org = org_admin
    emp, share, ds, twin = _shared_with_employee(app, admin, org, demo_files)
    got = emp.get(f"/api/received/{share['id']}").json()
    assert got["returns"] == {"table": "customers", "column": "customer_id", "ref_column": REF_COLUMN}
    customers = _download_csv_tables(emp, share["id"])["customers"]
    cols = list(customers.columns)
    assert cols[cols.index("customer_id") + 1] == REF_COLUMN and customers[REF_COLUMN].str.match(REF_RE).all()

    r = _submit(emp, share["id"], _results_csv(_scored(customers)))
    assert r.status_code == 201, r.text
    ret = r.json()
    assert ret["rows_accepted"] == len(customers) and sum(ret["rejected"].values()) == 0
    assert emp.get(f"/api/received/{share['id']}/returns").json()[0]["id"] == ret["id"]
    listed = admin.get(f"/api/orgs/{org}/returns").json()
    assert listed[0]["id"] == ret["id"] and listed[0]["relink"]["status"] == "none"

    caplog.set_level(logging.DEBUG)
    detail = _relink(app, admin, org, ret["id"], demo_files)
    assert detail["relink"]["status"] == "ready", detail["relink"]
    assert detail["relink"]["matched"] == len(customers)
    out = _relinked_csv(admin, org, ret["id"])
    original = pd.read_csv(io.BytesIO(dict(demo_files)["customers.csv"]), dtype=str)
    # row i of the twin is row i of the original: the score (row number) must land on the real key
    assert out["customer_id"].tolist() == [original["customer_id"][int(i)] for i in out["score"]]
    assert REF_COLUMN not in out.columns

    # the same admin may download again within the window; each download is logged
    assert admin.get(f"/api/orgs/{org}/returns/{ret['id']}/relinked.xlsx").status_code == 200
    with app.state.sessionmaker() as db:
        downloads = db.execute(select(AuditEvent).where(AuditEvent.action == "relink.downloaded")).scalars().all()
        assert [e.meta["download"] for e in downloads] == [1, 2]
        events = db.execute(select(AuditEvent)).scalars().all()
        kinds = {b.kind for b in db.execute(select(Blob)).scalars()}
    assert "relink_upload" not in kinds and "relinked" in kinds  # the uploaded originals are gone already

    # value-free: no real key, no pseudonym, no score in audit entries or logs
    secrets = set(original["customer_id"].head(50)) | set(original["national_id"].head(50)) | set(customers["customer_id"].head(50))
    blob = json.dumps([e.meta for e in events], ensure_ascii=False) + caplog.text
    assert not any(v in blob for v in secrets)

    # after the window: refused, then the sweep deletes the file
    with app.state.sessionmaker() as db:
        db.execute(update(Blob).where(Blob.kind == "relinked").values(expires_at=utcnow() - timedelta(seconds=1)))
        db.execute(update(Return).values(relink_expires_at=utcnow() - timedelta(seconds=1)))
        db.commit()
    assert admin.get(f"/api/orgs/{org}/returns/{ret['id']}/relinked.csv").status_code == 410
    with app.state.sessionmaker() as db:
        assert sweep(db, app.state.settings)["relink_files_deleted"] == 1
        assert db.execute(select(Blob).where(Blob.kind == "relinked")).first() is None
    assert admin.get(f"/api/orgs/{org}/returns/{ret['id']}").json()["relink"]["status"] == "expired"


# ---------------------------------------------------------------- rejection

def test_tampered_unknown_and_foreign_files_are_rejected(app, org_admin, demo_files):
    admin, org = org_admin
    emp, share, ds, twin = _shared_with_employee(app, admin, org, demo_files)
    customers = _download_csv_tables(emp, share["id"])["customers"]
    base = _scored(customers)
    n = len(base)

    # > 1% of rows with a swapped key (reference no longer matches): the whole file is refused
    swapped = base.copy()
    k = int(n * 0.02)
    swapped.loc[: k - 1, "customer_id"] = list(reversed(swapped["customer_id"][:k]))
    r = _submit(emp, share["id"], _results_csv(swapped))
    assert r.status_code == 422 and r.json()["code"] == "return_tampered"
    assert r.json()["rejected"]["ref_mismatch"] >= k - 1

    # <= 1% (here 1 row of 150): accepted, the bad row is counted and dropped
    assert n >= 100
    few = base.copy()
    few.loc[0, REF_COLUMN] = few.loc[1, REF_COLUMN]
    r = _submit(emp, share["id"], _results_csv(few))
    assert r.status_code == 201 and r.json()["rejected"]["ref_mismatch"] == 1 and r.json()["rows_accepted"] == n - 1

    # unknown keys: up to 5% accepted, more refused
    unknown = base.copy()
    unknown.loc[: int(n * 0.03), "customer_id"] = "ZZ-NOT-A-KEY"
    r = _submit(emp, share["id"], _results_csv(unknown))
    assert r.status_code == 201 and r.json()["rejected"]["unknown_key"] > 0
    unknown.loc[: int(n * 0.06), "customer_id"] = "ZZ-NOT-A-KEY"
    r = _submit(emp, share["id"], _results_csv(unknown))
    assert r.status_code == 422 and r.json()["code"] == "return_too_many_errors"

    # missing columns
    r = _submit(emp, share["id"], _results_csv(base.drop(columns=[REF_COLUMN])))
    assert r.json()["code"] == "return_missing_ref_column"
    r = _submit(emp, share["id"], _results_csv(base.drop(columns=["customer_id"])))
    assert r.json()["code"] == "return_missing_key_column"

    # a file made for another share (same twin) carries the other share's references
    member = next(m for m in admin.get(f"/api/orgs/{org}/members").json() if m["email"] == "employee@alwaha.example.com")
    other = admin.post(f"/api/orgs/{org}/twins/{twin['id']}/shares", json={"member_ids": [member["id"]]}).json()
    foreign = _download_csv_tables(emp, other["id"])["customers"]
    r = _submit(emp, share["id"], _results_csv(_scored(foreign)))
    assert r.status_code == 422 and r.json()["code"] == "return_tampered"

    with app.state.sessionmaker() as db:
        rejected = db.execute(select(AuditEvent).where(AuditEvent.action == "return.rejected")).scalars().all()
    assert {e.meta["reason"] for e in rejected} >= {"return_tampered", "return_too_many_errors"}


def test_relink_refuses_other_originals(app, org_admin, demo_files):
    admin, org = org_admin
    emp, share, ds, twin = _shared_with_employee(app, admin, org, demo_files)
    customers = _download_csv_tables(emp, share["id"])["customers"]
    ret = _submit(emp, share["id"], _results_csv(_scored(customers))).json()
    files = dict(demo_files)

    # a superset (one extra customer) is refused: extra values can change pseudonyms
    extra = files["customers.csv"].rstrip(b"\n") + b"\n" + files["customers.csv"].splitlines()[1].replace(b",", b"9,", 1) + b"\n"
    d = _relink(app, admin, org, ret["id"], [("customers.csv", extra), ("claims.csv", files["claims.csv"])])
    assert d["relink"]["status"] == "failed" and d["relink"]["error_code"] == "relink_rows_differ"
    # a renamed file is a different table
    d = _relink(app, admin, org, ret["id"], [("clients.csv", files["customers.csv"]), ("claims.csv", files["claims.csv"])])
    assert d["relink"]["error_code"] == "relink_tables_differ"
    # same shape, one key changed: the recomputed twin does not reproduce the shared one
    original = pd.read_csv(io.BytesIO(files["customers.csv"]), dtype=str, keep_default_na=False)
    original.loc[5, "customer_id"] = "999999"
    d = _relink(app, admin, org, ret["id"], [("customers.csv", original.to_csv(index=False).encode()),
                                              ("claims.csv", files["claims.csv"])])
    assert d["relink"]["error_code"] == "relink_mismatch"
    with app.state.sessionmaker() as db:
        assert db.execute(select(Blob).where(Blob.kind.in_(("relink_upload", "relinked")))).first() is None

    # after a key rotation the twin cannot be re-linked
    with app.state.sessionmaker() as db:
        db.execute(update(Organization).where(Organization.id == org).values(key_version=Organization.key_version + 1))
        db.commit()
    d = _relink(app, admin, org, ret["id"], demo_files)
    assert d["relink"]["error_code"] == "relink_key_rotated"


def test_superset_can_change_pseudonyms():
    """Why re-linking requires the exact original set: the collision step is order- and set-dependent."""
    import random

    rng = random.Random(1)
    base = sorted({str(rng.randint(100000, 999999)) for _ in range(3000)})
    extra = [str(rng.randint(100000, 999999)) for _ in range(500)]

    def run(values):
        p = Pseudonymizer(b"k" * 32)
        p.prepare({"PK:t.id": values})
        return {v: p.fake("PK:t.id", v) for v in values}

    a, b = run(base), run(base + extra)
    assert any(a[v] != b[v] for v in base)


# ---------------------------------------------------------------- cleaning consistency

def test_relink_reapplies_the_same_cleaning(app, org_admin, demo_files):
    """Pseudonyms are computed on cleaned values: the admin's raw originals (extra spaces, mixed mobile
    formats, duplicate rows) must be cleaned the same way before recomputing, and still re-link fully."""
    admin, org = org_admin
    files = dict(demo_files)
    cust = pd.read_csv(io.BytesIO(files["customers.csv"]), dtype=str, keep_default_na=False)
    cust["full_name"] = cust["full_name"].map(lambda v: "  " + v.replace(" ", "   ") + " ")
    cust["mobile"] = [("+966" + m[1:]) if i % 3 == 0 else m for i, m in enumerate(cust["mobile"])]
    cust = pd.concat([cust, cust.head(25)], ignore_index=True)  # exact duplicate rows
    messy = [("customers.csv", cust.to_csv(index=False).encode("utf-8")), ("claims.csv", files["claims.csv"])]

    emp, _ = invite_and_join(app, admin, org, "employee@alwaha.example.com")
    ds = ready_dataset(app, admin, org, messy)
    assert ds["summary"]["cleaning"]["report"]["applied"]["dedupe"]["total"] == 25
    assert admin.post(f"/api/orgs/{org}/datasets/{ds['id']}/clean", json={"phones": True}).status_code == 202
    work(app)
    twin = masked_twin(app, admin, org, ds["id"], apply_fix="auto")
    assert twin["verdict"] == "PASS"
    member = next(m for m in admin.get(f"/api/orgs/{org}/members").json() if m["email"] == "employee@alwaha.example.com")
    share = admin.post(f"/api/orgs/{org}/twins/{twin['id']}/shares", json={"member_ids": [member["id"]]}).json()
    customers = _download_csv_tables(emp, share["id"])["customers"]
    assert len(customers) == len(cust) - 25
    ret = _submit(emp, share["id"], _results_csv(_scored(customers))).json()

    d = _relink(app, admin, org, ret["id"], messy)
    assert d["relink"]["status"] == "ready" and d["relink"]["matched"] == len(customers), d["relink"]
    out = _relinked_csv(admin, org, ret["id"])
    deduped = cust.drop_duplicates().reset_index(drop=True)
    assert out["customer_id"].tolist() == [deduped["customer_id"][int(i)] for i in out["score"]]


# ---------------------------------------------------------------- Excel safety

def _twin_with(keys):
    return {"t": pd.DataFrame({"id": keys, "x": ["a"] * len(keys)})}


def _excel_csv_cell(v: str) -> str:
    """What Excel writes back after opening a CSV and saving it: number-like text becomes a number;
    leading zeros are lost; 12+ digits are shown (and saved) in scientific notation."""
    if re.fullmatch(r"\d+", v):
        n = int(v)
        return f"{n:.5E}".replace("E+0", "E+") if len(str(n)) >= 12 else str(n)
    if re.fullmatch(r"\d+\.\d+", v):
        return repr(float(v))
    return v


def test_refs_and_keys_survive_excel_round_trips():
    key = b"r" * 32
    keys = ["751233", "007431", "1234567890123", "4402917765", "000123456789012"]
    twin = _twin_with(keys)
    link = {"table": "t", "column": "id"}
    refs = [make_ref(key, k) for k in keys]
    assert all(REF_RE.match(r) for r in refs)
    assert not any(re.fullmatch(r"[\d.eE+-]+", r[3:]) for r in refs)  # never digits-only or e-notation
    results = pd.DataFrame({"id": keys, REF_COLUMN: refs, "score": ["0.5", "1", "2", "3", "4"]})

    # 1) through openpyxl: export (all text cells), open, save, read back
    from nazeer_api.returns import add_refs
    exported = export_xlsx(add_refs(twin, link, key))
    wb = load_workbook(io.BytesIO(exported))
    buf = io.BytesIO()
    wb.save(buf)
    ws = load_workbook(io.BytesIO(buf.getvalue())).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:2] == ("id", REF_COLUMN)
    assert [r[0] for r in rows[1:]] == keys and [r[1] for r in rows[1:]] == refs

    # a recipient's own workbook where Excel stored keys as numbers: still accepted, keys restored
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.append(["id", REF_COLUMN, "score"])
    for k, r in zip(keys, refs):
        ws2.append([int(k) if len(k) < 16 else k, r, 1])
    buf2 = io.BytesIO()
    wb2.save(buf2)
    out, stats = check_return(twin, link, key, [("r.xlsx", buf2.getvalue())])
    assert out["id"].tolist() == keys and stats["rows_accepted"] == 5

    # 2) through a CSV opened and saved by Excel
    csv_text = io.StringIO()
    excelled = results.map(lambda v: _excel_csv_cell(str(v)))
    assert excelled[REF_COLUMN].tolist() == refs                  # references untouched
    assert excelled["id"].tolist() != keys                        # keys were mangled...
    excelled.to_csv(csv_text, index=False)
    out, stats = check_return(twin, link, key, [("r.csv", csv_text.getvalue().encode("utf-8"))])
    assert out["id"].tolist() == keys and stats["excel_repaired"] >= 3  # ...and restored exactly

    # the twin's own CSV export keeps keys as text (Excel-only damage is repaired on return)
    with zipfile.ZipFile(io.BytesIO(export_csv_zip(add_refs(twin, link, key)))) as z:
        back = pd.read_csv(z.open("t.csv"), dtype=str, encoding="utf-8-sig")
    assert back["id"].tolist() == keys and back[REF_COLUMN].tolist() == refs


def test_reference_does_not_cover_result_values():
    """Documented limit: nazeer_ref binds the key, not the recipient's result values."""
    key = b"r" * 32
    twin = _twin_with(["751233", "751234"])
    link = {"table": "t", "column": "id"}
    df = pd.DataFrame({"id": ["751233", "751234"], REF_COLUMN: [make_ref(key, "751233"), make_ref(key, "751234")],
                       "score": ["0.1", "0.2"]})
    df.loc[0, "score"] = "0.9"  # edited result value: still accepted
    out, _ = check_return(twin, link, key, [("r.csv", df.to_csv(index=False).encode())])
    assert out["score"].tolist() == ["0.9", "0.2"]
    with pytest.raises(ReturnError):
        df.loc[0, "id"] = "751234"  # an edited key: caught
        check_return(twin, link, key, [("r.csv", df.to_csv(index=False).encode())])


# ---------------------------------------------------------------- roles and tenancy

def test_roles_and_tenancy_for_returns(app, org_admin, demo_files):
    from tests_api.conftest import signup

    admin, org = org_admin
    emp, share, ds, twin = _shared_with_employee(app, admin, org, demo_files)
    customers = _download_csv_tables(emp, share["id"])["customers"]
    ret = _submit(emp, share["id"], _results_csv(_scored(customers))).json()

    dm, _ = invite_and_join(app, admin, org, "dm@alwaha.example.com", data_manager=True)
    assert dm.get(f"/api/orgs/{org}/returns").status_code == 200
    r = dm.post(f"/api/orgs/{org}/returns/{ret['id']}/relink", files=[("files", ("customers.csv", b"a\n1", "text/csv"))])
    assert r.status_code == 403
    assert emp.get(f"/api/orgs/{org}/returns").status_code in (403, 404)

    second_admin, _ = invite_and_join(app, admin, org, "admin2@alwaha.example.com", role="admin")
    _relink(app, admin, org, ret["id"], demo_files)
    assert second_admin.get(f"/api/orgs/{org}/returns/{ret['id']}/relinked.csv").json()["code"] == "relink_not_yours"

    other, me = signup(app, "boss@other.example.com", org="منشأة أخرى")
    other_org = me["memberships"][0]["org_id"]
    assert other.get(f"/api/orgs/{org}/returns/{ret['id']}").status_code == 404
    assert other.get(f"/api/orgs/{other_org}/returns/{ret['id']}").status_code == 404
    assert other.post(f"/api/received/{share['id']}/returns",
                      files=[("files", ("r.csv", b"a\n1", "text/csv"))]).status_code == 404


def test_synthetic_twins_take_no_returns(app, org_admin):
    from nazeer_api.returns import link_of

    class T:
        mode, report = "synthetic", {"generation": {"link": {"table": "t", "column": "id"}}}

    assert link_of(T()) is None
