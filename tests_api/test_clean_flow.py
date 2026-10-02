"""The common real case: an already-clean file goes through the whole flow with ZERO cleaning changes:
upload -> detection -> twin -> report -> share -> return (a subset with predictions) -> re-link."""
import io
import zipfile
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from tests_api.conftest import invite_and_join
from tests_api.test_p2 import masked_twin, org_admin, ready_dataset, work  # noqa: F401
from tests_api.test_p5_tokens import to_xlsx

SAMPLE = Path(__file__).resolve().parents[1] / "web" / "public" / "samples" / "clinic_appointments_clean.csv"


def test_clean_file_full_flow_without_any_cleaning(app, org_admin):
    admin, org = org_admin
    data = SAMPLE.read_bytes()
    ds = ready_dataset(app, admin, org, [(SAMPLE.name, data)])
    c = ds["summary"]["cleaning"]
    assert c["decision"] == "not_needed" and c["recommended"] == 0             # «البيانات نظيفة»
    assert sum(v["total"] for v in c["report"]["applied"].values()) == 0
    assert ds["summary"]["total_rows"] == 1000

    twin = masked_twin(app, admin, org, ds["id"])
    assert twin["verdict"] == "PASS", twin["failed_checks"]
    assert twin["summary"]["cleaned"] == {} and twin["summary"]["cleaning_decision"] == "not_needed"
    assert admin.get(f"/api/orgs/{org}/twins/{twin['id']}/report.pdf").status_code == 200

    emp, _ = invite_and_join(app, admin, org, "analyst@alwaha.example.com")
    member = next(m for m in admin.get(f"/api/orgs/{org}/members").json() if m["email"] == "analyst@alwaha.example.com")
    share = admin.post(f"/api/orgs/{org}/twins/{twin['id']}/shares", json={"member_ids": [member["id"]]}).json()
    url = emp.post(f"/api/received/{share['id']}/download", json={"format": "xlsx"}).json()["url"]
    ws = load_workbook(io.BytesIO(emp.get(url).content)).active
    rows = list(ws.iter_rows(values_only=True))
    shared = pd.DataFrame(rows[1:], columns=rows[0])
    sub = shared.head(200).copy()
    sub["احتمال_عدم_الحضور"] = [f"{(i % 10) / 10:.1f}" for i in range(200)]          # the model's predictions
    r = emp.post(f"/api/received/{share['id']}/returns",
                 files=[("files", ("predictions.xlsx", to_xlsx(sub), "application/octet-stream"))]).json()
    assert r["report"]["integrity"] == 1.0 and r["report"]["coverage"] == 0.2

    r2 = admin.post(f"/api/orgs/{org}/returns/{r['id']}/relink", files=[("files", (SAMPLE.name, data, "text/csv"))])
    assert r2.status_code == 202
    work(app)
    out = admin.get(f"/api/orgs/{org}/returns/{r['id']}/relinked.csv")
    with zipfile.ZipFile(io.BytesIO(out.content)) as z:
        linked = pd.read_csv(z.open(z.namelist()[0]), dtype=str, keep_default_na=False, encoding="utf-8-sig")
    original = pd.read_csv(io.BytesIO(data), dtype=str, keep_default_na=False)
    assert len(linked) == 200 and set(linked["حالة_الربط"]) == {"مرتبط"}
    by_id = original.set_index("رقم_الموعد")
    assert (linked["رقم_الهوية"] == by_id.loc[linked["رقم_الموعد"], "رقم_الهوية"].values).all()
    assert "احتمال_عدم_الحضور" in linked.columns
